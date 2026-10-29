import csv
import io
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from .schemas import schema_fields

MAX_IMPORT_ROWS = 5000
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}


class ImportFileError(ValueError):
    pass


def _cell_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None:
        return ""
    return value


def _normalize_rows(records, column_mapping, canonical_fields):
    normalized = []
    for row_number, raw in records:
        mapped = {}
        for field in canonical_fields:
            source_header = column_mapping.get(field, field)
            mapped[field] = _cell_value(raw.get(source_header, ""))
        normalized.append((row_number, raw, mapped))
    return normalized


def _parse_csv(data, column_mapping, canonical_fields):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError("CSV must use UTF-8 encoding.") from exc
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ImportFileError("CSV header row is missing.")
    headers = [header.strip() for header in reader.fieldnames]
    if any(not header for header in headers):
        raise ImportFileError("CSV contains an empty header.")
    if len(headers) != len(set(headers)):
        raise ImportFileError("CSV contains duplicate headers.")
    reader.fieldnames = headers
    records = []
    for row_number, raw in enumerate(reader, start=2):
        if len(records) >= MAX_IMPORT_ROWS:
            raise ImportFileError(f"Import exceeds {MAX_IMPORT_ROWS} rows.")
        clean = {
            key: _cell_value(value) for key, value in raw.items() if key is not None
        }
        if any(str(value).strip() for value in clean.values()):
            records.append((row_number, clean))
    return headers, _normalize_rows(records, column_mapping, canonical_fields)


def _parse_xlsx(data, column_mapping, canonical_fields):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > (
                MAX_XLSX_UNCOMPRESSED_BYTES
            ):
                raise ImportFileError("XLSX expands beyond the safety limit.")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("File is not a valid XLSX workbook.") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as exc:
        raise ImportFileError("Workbook is empty.") from exc
    headers = [
        str(value).strip() if value is not None else "" for value in header_values
    ]
    while headers and not headers[-1]:
        headers.pop()
    if any(not header for header in headers):
        raise ImportFileError("Workbook contains an empty header.")
    if len(headers) != len(set(headers)):
        raise ImportFileError("Workbook contains duplicate headers.")
    records = []
    for row_number, values in enumerate(rows, start=2):
        if len(records) >= MAX_IMPORT_ROWS:
            raise ImportFileError(f"Import exceeds {MAX_IMPORT_ROWS} rows.")
        clean = {
            headers[index]: _cell_value(value)
            for index, value in enumerate(values)
            if index < len(headers)
        }
        if any(str(value).strip() for value in clean.values()):
            records.append((row_number, clean))
    return headers, _normalize_rows(records, column_mapping, canonical_fields)


def parse_import_file(batch):
    extension = f".{batch.original_filename.rsplit('.', 1)[-1].lower()}"
    if extension not in ALLOWED_EXTENSIONS:
        raise ImportFileError("Only .csv and .xlsx files are accepted.")
    if batch.file_size > MAX_FILE_BYTES:
        raise ImportFileError("File exceeds the 10 MB upload limit.")

    with batch.file.open("rb") as file:
        data = file.read()
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("File exceeds the 10 MB upload limit.")

    fields = schema_fields(batch.import_type)
    invalid_mapping = set(batch.column_mapping) - set(fields)
    if invalid_mapping:
        raise ImportFileError(
            "Unknown mapping fields: " + ", ".join(sorted(invalid_mapping))
        )

    if extension == ".csv":
        headers, rows = _parse_csv(data, batch.column_mapping, fields)
    else:
        headers, rows = _parse_xlsx(data, batch.column_mapping, fields)

    missing = {
        field
        for field in fields
        if batch.column_mapping.get(field, field) not in headers
    }
    if missing:
        raise ImportFileError(
            "Missing required column(s): " + ", ".join(sorted(missing))
        )

    if len(rows) > MAX_IMPORT_ROWS:
        raise ImportFileError(f"Import exceeds {MAX_IMPORT_ROWS} rows.")
    if not rows:
        raise ImportFileError("Import contains no data rows.")
    return headers, rows

import csv
import json
from io import BytesIO, StringIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.context import tenant_context
from data_imports.models import ImportBatch
from data_imports.schemas import template_headers
from data_imports.tasks import cleanup_expired_import_files
from fleet.models import Accessory, Asset, Location
from tenants.models import Membership, Plan, Role
from tests_support.factories import (
    make_equipment_model,
    make_location,
    make_owner,
    make_plan,
    make_tenant,
)


class ImportApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tenant = make_tenant("import-api", plan=make_plan("import-api-plan"))
        cls.owner = make_owner(cls.tenant, "import-owner")
        cls.model = make_equipment_model(cls.tenant, "IMPORT-MODEL")
        cls.location = make_location(cls.tenant, "IMPORT-WH")

    def setUp(self):
        self.client.force_login(self.owner)

    def csv_file(self, import_type, rows, headers=None):
        headers = headers or template_headers(import_type)
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
        return SimpleUploadedFile(
            f"{import_type}.csv",
            output.getvalue().encode("utf-8-sig"),
            content_type="text/csv",
        )

    def upload(
        self,
        import_type,
        rows,
        *,
        duplicate_policy="error",
        mapping=None,
        mapping_name="",
        headers=None,
    ):
        return self.client.post(
            "/api/v1/imports/",
            {
                "file": self.csv_file(import_type, rows, headers),
                "import_type": import_type,
                "duplicate_policy": duplicate_policy,
                "column_mapping": mapping or {},
                "mapping_name": mapping_name,
            },
            HTTP_X_TENANT=self.tenant.slug,
        )

    def get_batch(self, batch_id):
        with tenant_context(self.tenant):
            return ImportBatch.objects.get(pk=batch_id)

    def test_xlsx_and_csv_templates_are_available(self):
        xlsx = self.client.get(
            "/api/v1/import-templates/asset/?file_format=xlsx",
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(xlsx.status_code, 200)
        workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
        self.assertEqual(
            [cell.value for cell in next(workbook["Data"].iter_rows())],
            template_headers("asset"),
        )
        csv_response = self.client.get(
            "/api/v1/import-templates/asset/?file_format=csv",
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(csv_response.status_code, 200)
        self.assertTrue(csv_response.content.startswith(b"\xef\xbb\xbf"))

    def test_xlsx_upload_is_validated_and_committed(self):
        workbook = Workbook()
        sheet = workbook.active
        headers = template_headers("location")
        sheet.append(headers)
        sheet.append(["XLSX-WH", "XLSX Warehouse", "warehouse", "4 Road", "true"])
        output = BytesIO()
        workbook.save(output)
        response = self.client.post(
            "/api/v1/imports/",
            {
                "file": SimpleUploadedFile(
                    "locations.xlsx",
                    output.getvalue(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                ),
                "import_type": "location",
                "duplicate_policy": "error",
            },
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(response.status_code, 202)
        batch = self.get_batch(response.json()["id"])
        self.assertEqual(batch.error_rows, 0)
        commit = self.client.post(
            f"/api/v1/imports/{batch.pk}/commit/",
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(commit.status_code, 200)
        with tenant_context(self.tenant):
            self.assertTrue(Location.objects.filter(code="XLSX-WH").exists())

    def test_location_preview_commit_idempotency_and_error_report(self):
        response = self.upload(
            "location",
            [
                {
                    "code": "NEW-WH",
                    "name": "New Warehouse",
                    "location_type": "warehouse",
                    "address": "2 Test Road",
                    "is_active": "true",
                }
            ],
        )
        self.assertEqual(response.status_code, 202)
        batch = self.get_batch(response.json()["id"])
        self.assertEqual(batch.status, ImportBatch.Status.PREVIEW_READY)
        self.assertEqual(batch.error_rows, 0)
        self.assertTrue(batch.can_commit)

        committed = self.client.post(
            f"/api/v1/imports/{batch.pk}/commit/",
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(committed.status_code, 200)
        self.assertEqual(committed.json()["status"], "completed")
        with tenant_context(self.tenant):
            self.assertTrue(Location.objects.filter(code="NEW-WH").exists())
        repeated = self.client.post(
            f"/api/v1/imports/{batch.pk}/commit/",
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(repeated.status_code, 400)

        report = self.client.get(
            f"/api/v1/imports/{batch.pk}/error-report/",
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(report.status_code, 200)
        self.assertIn("row_number", report.content.decode("utf-8-sig"))

    def test_preview_batch_can_be_cancelled(self):
        response = self.upload(
            "location",
            [
                {
                    "code": "CANCEL-WH",
                    "name": "Cancel Warehouse",
                    "location_type": "warehouse",
                    "address": "",
                    "is_active": "true",
                }
            ],
        )
        self.assertEqual(response.status_code, 202)
        batch_id = response.json()["id"]

        cancelled = self.client.post(
            f"/api/v1/imports/{batch_id}/cancel/",
            HTTP_X_TENANT=self.tenant.slug,
        )

        self.assertEqual(cancelled.status_code, 200)
        self.assertEqual(cancelled.json()["status"], "cancelled")
        self.assertEqual(self.get_batch(batch_id).status, ImportBatch.Status.CANCELLED)

    def test_custom_mapping_is_saved_and_used(self):
        headers = ["Location ID", "Location Name", "Type", "Address", "Active"]
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerow(["MAP-WH", "Mapped Warehouse", "warehouse", "3 Road", "true"])
        response = self.client.post(
            "/api/v1/imports/",
            {
                "file": SimpleUploadedFile(
                    "mapped.csv",
                    output.getvalue().encode(),
                    content_type="text/csv",
                ),
                "import_type": "location",
                "duplicate_policy": "error",
                "column_mapping": json.dumps(
                    {
                        "code": "Location ID",
                        "name": "Location Name",
                        "location_type": "Type",
                        "address": "Address",
                        "is_active": "Active",
                    }
                ),
                "mapping_name": "Legacy location export",
            },
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(response.status_code, 202)
        batch = self.get_batch(response.json()["id"])
        self.assertEqual(batch.status, ImportBatch.Status.PREVIEW_READY)
        self.assertEqual(batch.error_rows, 0)

    def test_formula_content_is_neutralized_in_error_report(self):
        response = self.upload(
            "location",
            [
                {
                    "code": "",
                    "name": '=HYPERLINK("https://example.invalid")',
                    "location_type": "warehouse",
                    "address": "",
                    "is_active": "true",
                }
            ],
        )
        batch = self.get_batch(response.json()["id"])
        self.assertEqual(batch.error_rows, 1)
        report = self.client.get(
            f"/api/v1/imports/{batch.pk}/error-report/",
            HTTP_X_TENANT=self.tenant.slug,
        )
        body = report.content.decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", body)

    def test_repeated_file_is_flagged_but_can_be_reuploaded(self):
        row = {
            "code": "DUP-WH",
            "name": "Duplicate Warehouse",
            "location_type": "warehouse",
            "address": "",
            "is_active": "true",
        }
        first = self.upload("location", [row])
        self.assertEqual(first.status_code, 202)
        second = self.upload("location", [row], duplicate_policy="skip")
        self.assertEqual(second.status_code, 202)
        batch = self.get_batch(second.json()["id"])
        self.assertEqual(batch.summary["duplicate_file_of"], first.json()["id"])

    def test_invalid_extension_and_large_file_are_rejected(self):
        invalid = self.client.post(
            "/api/v1/imports/",
            {
                "file": SimpleUploadedFile("assets.xls", b"legacy"),
                "import_type": "asset",
            },
            HTTP_X_TENANT=self.tenant.slug,
        )
        self.assertEqual(invalid.status_code, 400)

        rows = [
            {
                "code": f"BULK-{index}",
                "name": f"Bulk {index}",
                "location_type": "warehouse",
                "address": "",
                "is_active": "true",
            }
            for index in range(5001)
        ]
        response = self.upload("location", rows)
        self.assertEqual(response.status_code, 202)
        batch = self.get_batch(response.json()["id"])
        self.assertEqual(batch.status, ImportBatch.Status.FAILED)
        self.assertIn("5000", batch.failure_message)

    def test_expired_import_file_content_can_be_cleaned(self):
        response = self.upload(
            "location",
            [
                {
                    "code": "EXPIRY-WH",
                    "name": "Expiry Warehouse",
                    "location_type": "warehouse",
                    "address": "",
                    "is_active": "true",
                }
            ],
        )
        batch = self.get_batch(response.json()["id"])
        self.assertTrue(batch.file.name)
        with tenant_context(self.tenant):
            ImportBatch.objects.filter(pk=batch.pk).update(
                file_expires_at=timezone.now() - timezone.timedelta(days=1)
            )
        cleanup_expired_import_files()
        batch.refresh_from_db()
        self.assertFalse(batch.file.name)

    def test_cross_tenant_batches_are_hidden(self):
        response = self.upload(
            "location",
            [
                {
                    "code": "TENANT-A-WH",
                    "name": "Tenant A Warehouse",
                    "location_type": "warehouse",
                    "address": "",
                    "is_active": "true",
                }
            ],
        )
        batch_id = response.json()["id"]
        other = make_tenant("import-other", plan=make_plan("import-other-plan"))
        other_owner = make_owner(other, "import-other-owner")
        self.client.force_login(other_owner)
        response = self.client.get(
            f"/api/v1/imports/{batch_id}/",
            HTTP_X_TENANT=other.slug,
        )
        self.assertEqual(response.status_code, 404)


class AssetAndAccessoryImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.plan = make_plan("asset-import-plan", max_assets=1)
        cls.tenant = make_tenant("asset-import", plan=cls.plan)
        cls.owner = make_owner(cls.tenant, "asset-import-owner")
        cls.model = make_equipment_model(cls.tenant, "ASSET-IMPORT-MODEL")
        cls.location = make_location(cls.tenant, "ASSET-IMPORT-WH")

    def setUp(self):
        self.client.force_login(self.owner)
        self.headers = {"HTTP_X_TENANT": self.tenant.slug}

    def upload_csv(self, import_type, rows, duplicate_policy="error"):
        output = StringIO()
        headers = template_headers(import_type)
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
        return self.client.post(
            "/api/v1/imports/",
            {
                "file": SimpleUploadedFile(
                    f"{import_type}.csv",
                    output.getvalue().encode("utf-8-sig"),
                    content_type="text/csv",
                ),
                "import_type": import_type,
                "duplicate_policy": duplicate_policy,
            },
            **self.headers,
        )

    def batch(self, batch_id):
        with tenant_context(self.tenant):
            return ImportBatch.objects.get(pk=batch_id)

    def test_asset_quota_blocks_preview_commit(self):
        response = self.upload_csv(
            "asset",
            [
                {
                    "asset_number": "IMP-1",
                    "model_code": self.model.code,
                    "location_code": self.location.code,
                    "serial_number": "S-1",
                    "current_hours": "1",
                },
                {
                    "asset_number": "IMP-2",
                    "model_code": self.model.code,
                    "location_code": self.location.code,
                    "serial_number": "S-2",
                    "current_hours": "1",
                },
            ],
        )
        batch = self.batch(response.json()["id"])
        self.assertEqual(batch.error_rows, 1)
        self.assertFalse(batch.can_commit)
        commit = self.client.post(f"/api/v1/imports/{batch.pk}/commit/", **self.headers)
        self.assertEqual(commit.status_code, 400)

    def test_asset_duplicate_update_and_undo_safe_state(self):
        with tenant_context(self.tenant):
            asset = Asset.objects.create(
                model=self.model,
                asset_number="EXISTING-1",
                serial_number="OLD-SERIAL",
                notes="old",
            )
        response = self.upload_csv(
            "asset",
            [
                {
                    "asset_number": "EXISTING-1",
                    "model_code": self.model.code,
                    "location_code": self.location.code,
                    "serial_number": "OLD-SERIAL",
                    "current_hours": "0",
                    "notes": "updated by import",
                }
            ],
            duplicate_policy="update",
        )
        batch = self.batch(response.json()["id"])
        commit = self.client.post(f"/api/v1/imports/{batch.pk}/commit/", **self.headers)
        self.assertEqual(commit.status_code, 200)
        with tenant_context(self.tenant):
            asset.refresh_from_db()
            self.assertEqual(asset.notes, "updated by import")
            self.assertFalse(batch.can_undo)

    def test_accessory_import_resolves_compatible_models(self):
        rows = [
            {
                "code": "BUCKET-1",
                "name": "Heavy duty bucket",
                "accessory_type": "bucket",
                "serial_number": "B-1",
                "location_code": self.location.code,
                "compatible_model_codes": self.model.code,
                "status": "available",
                "notes": "",
            }
        ]
        response = self.upload_csv("accessory", rows)
        batch = self.batch(response.json()["id"])
        self.assertEqual(batch.error_rows, 0)
        commit = self.client.post(f"/api/v1/imports/{batch.pk}/commit/", **self.headers)
        self.assertEqual(commit.status_code, 200)
        with tenant_context(self.tenant):
            accessory = Accessory.objects.get(code="BUCKET-1")
            self.assertEqual(list(accessory.compatible_models.all()), [self.model])

    def test_meter_reading_import_updates_asset_but_cannot_be_undone(self):
        with tenant_context(self.tenant):
            asset = Asset.objects.create(
                model=self.model,
                asset_number="METER-1",
                current_hours=5,
            )
        response = self.upload_csv(
            "meter_reading",
            [
                {
                    "asset_number": asset.asset_number,
                    "hours": "6.5",
                    "kilometers": "",
                    "recorded_at": "2026-10-29T08:00:00+08:00",
                    "source": "import",
                    "notes": "Opening reading",
                }
            ],
        )
        batch = self.batch(response.json()["id"])
        commit = self.client.post(f"/api/v1/imports/{batch.pk}/commit/", **self.headers)
        self.assertEqual(commit.status_code, 200)
        with tenant_context(self.tenant):
            asset.refresh_from_db()
            self.assertEqual(asset.current_hours, 6.5)
        self.assertFalse(self.batch(batch.pk).can_undo)

    def test_location_import_can_be_undone_before_use(self):
        response = self.upload_csv(
            "location",
            [
                {
                    "code": "UNDO-WH",
                    "name": "Undo Warehouse",
                    "location_type": "warehouse",
                    "address": "",
                    "is_active": "true",
                }
            ],
        )
        batch = self.batch(response.json()["id"])
        self.client.post(f"/api/v1/imports/{batch.pk}/commit/", **self.headers)
        undo = self.client.post(f"/api/v1/imports/{batch.pk}/undo/", **self.headers)
        self.assertEqual(undo.status_code, 200)
        self.assertEqual(undo.json()["status"], "undone")
        with tenant_context(self.tenant):
            self.assertFalse(Location.objects.filter(code="UNDO-WH").exists())

    def test_imported_asset_is_retired_when_batch_is_undone(self):
        response = self.upload_csv(
            "asset",
            [
                {
                    "asset_number": "UNDO-ASSET",
                    "model_code": self.model.code,
                    "location_code": self.location.code,
                    "serial_number": "UNDO-ASSET-SERIAL",
                    "current_hours": "2",
                }
            ],
        )
        batch = self.batch(response.json()["id"])
        self.client.post(f"/api/v1/imports/{batch.pk}/commit/", **self.headers)
        with tenant_context(self.tenant):
            self.assertTrue(ImportBatch.objects.get(pk=batch.pk).can_undo)
        undo = self.client.post(f"/api/v1/imports/{batch.pk}/undo/", **self.headers)
        self.assertEqual(undo.status_code, 200)
        with tenant_context(self.tenant):
            self.assertEqual(
                Asset.objects.get(asset_number="UNDO-ASSET").status,
                Asset.Status.RETIRED,
            )

    def test_plan_feature_and_role_are_enforced(self):
        no_feature_plan = Plan.objects.create(
            code="no-import-plan", name="No import", max_assets=10
        )
        no_feature_tenant = make_tenant("no-import-tenant", plan=no_feature_plan)
        owner = make_owner(no_feature_tenant, "no-import-owner")
        self.client.force_login(owner)
        denied = self.client.get(
            "/api/v1/imports/", HTTP_X_TENANT=no_feature_tenant.slug
        )
        self.assertEqual(denied.status_code, 403)

        viewer = make_owner(self.tenant, "import-viewer")
        Membership.objects.filter(user=viewer, tenant=self.tenant).update(
            role=Role.VIEWER
        )
        self.client.force_login(viewer)
        response = self.upload_csv(
            "location",
            [
                {
                    "code": "VIEWER-WH",
                    "name": "Viewer Warehouse",
                    "location_type": "warehouse",
                    "address": "",
                    "is_active": "true",
                }
            ],
        )
        self.assertEqual(response.status_code, 403)

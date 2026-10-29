import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ImportBatch


IMPORT_SCHEMAS = {
    ImportBatch.Type.LOCATION: {
        "title": "Locations",
        "fields": [
            ("code", "Location code", True, "e.g. WH-KT"),
            ("name", "Location name", True, "Main warehouse"),
            ("location_type", "Type", True, "warehouse / yard / site / repair / other"),
            ("address", "Address", False, ""),
            ("is_active", "Active", True, "true / false"),
        ],
    },
    ImportBatch.Type.EQUIPMENT_MODEL: {
        "title": "Equipment models",
        "fields": [
            ("code", "Model code", True, "e.g. CAT-320"),
            ("name", "Model name", True, "CAT 320 Excavator"),
            ("category_code", "Category code", True, "Must already exist"),
            ("manufacturer", "Manufacturer", False, ""),
            ("model_number", "Manufacturer model", False, ""),
            ("specifications_json", "Specifications JSON", False, '{"bucket_m3": 1.4}'),
            ("is_active", "Active", True, "true / false"),
        ],
    },
    ImportBatch.Type.ASSET: {
        "title": "Assets",
        "fields": [
            ("asset_number", "Asset number", True, "Unique per tenant"),
            ("model_code", "Equipment model code", True, "Must already exist"),
            ("location_code", "Location code", False, "Must already exist"),
            ("serial_number", "Serial number", False, ""),
            ("registration_number", "Registration number", False, ""),
            ("qr_code", "QR code", False, ""),
            ("current_hours", "Current hours", False, "Only forward readings"),
            ("current_kilometers", "Current kilometres", False, ""),
            ("acquired_on", "Acquired date", False, "YYYY-MM-DD"),
            ("purchase_cost", "Purchase cost", False, ""),
            ("notes", "Notes", False, ""),
        ],
    },
    ImportBatch.Type.ACCESSORY: {
        "title": "Accessories",
        "fields": [
            ("code", "Accessory code", True, "Unique per tenant"),
            ("name", "Accessory name", True, ""),
            ("accessory_type", "Accessory type", False, "bucket / attachment / cable"),
            ("serial_number", "Serial number", False, ""),
            ("location_code", "Location code", False, "Must already exist"),
            (
                "compatible_model_codes",
                "Compatible model codes",
                False,
                "Comma separated",
            ),
            ("status", "Status", True, "available / in_use / maintenance / retired"),
            ("notes", "Notes", False, ""),
        ],
    },
    ImportBatch.Type.METER_READING: {
        "title": "Opening meter readings",
        "fields": [
            ("asset_number", "Asset number", True, "Must already exist"),
            ("hours", "Hours", False, "At least hours or kilometres"),
            ("kilometers", "Kilometres", False, "At least hours or kilometres"),
            ("recorded_at", "Recorded at", True, "ISO date/time"),
            ("source", "Source", True, "manual / telematics / import"),
            ("notes", "Notes", False, ""),
        ],
    },
}


def schema_fields(import_type):
    return [field[0] for field in IMPORT_SCHEMAS[import_type]["fields"]]


def template_headers(import_type):
    return [field[0] for field in IMPORT_SCHEMAS[import_type]["fields"]]


def build_xlsx_template(import_type):
    schema = IMPORT_SCHEMAS[import_type]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    header_fill = PatternFill("solid", fgColor="0C6B5F")
    for column, (field, label, required, help_text) in enumerate(
        schema["fields"], start=1
    ):
        cell = sheet.cell(row=1, column=column, value=field)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = max(
            16, min(38, len(help_text) + 4)
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(schema['fields']))}1"

    notes = workbook.create_sheet("Notes")
    notes.append(["Field", "Required", "Guidance"])
    for field, label, required, help_text in schema["fields"]:
        notes.append([field, "Yes" if required else "No", help_text or label])
    notes.column_dimensions["A"].width = 30
    notes.column_dimensions["B"].width = 12
    notes.column_dimensions["C"].width = 60
    notes["A1"].font = Font(bold=True)
    notes["B1"].font = Font(bold=True)
    notes["C1"].font = Font(bold=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_csv_template(import_type):
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(template_headers(import_type))
    return output.getvalue().encode("utf-8-sig")
